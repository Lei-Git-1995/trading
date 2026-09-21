"""Excel 选股报告：输出 output/YYYY-MM-DD.xlsx"""
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill('solid', fgColor='4472C4')
HEADER_FONT = Font(color='FFFFFF', bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=12, color='1F4E79')
GOOD_FILL = PatternFill('solid', fgColor='C6EFCE')
WARN_FILL = PatternFill('solid', fgColor='FFEB9C')


def _write_table(ws, start_row, headers, rows, widths=None):
    """写入表头 + 数据行，返回下一空行"""
    for col, head in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=head)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER

    for r, row in enumerate(rows, start_row + 1):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal='center')

    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return start_row + 1 + len(rows)


def generate(stocks: list, target_date: str, path: str, config) -> str:
    """生成包含多工作表的 Excel 报告，返回文件路径"""
    wb = Workbook()

    _write_summary(wb, stocks, target_date, config)
    _write_picks(wb, stocks)
    _write_top5(wb, stocks[:5])
    _write_history(wb, stocks)

    wb.save(path)
    return path


def _write_summary(wb, stocks, target_date, config):
    ws = wb.active
    ws.title = '报告说明'
    ws.cell(1, 1, '短线选股报告').font = Font(bold=True, size=14)
    ws['A2'] = f'行情数据日期：{target_date}（最近交易日快照）'
    ws['A3'] = f'生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'
    ws['A4'] = f'符合条件的股票：{len(stocks)} 只'
    ws['A6'] = '筛选条件'
    for i, cond in enumerate(
        [
            f'换手率 > {config.turnover_min}%',
            f'涨幅在 {config.change_min}% 到 {config.change_max}% 之间',
            f'量能比前两日平均放量 > {config.volume_ratio}倍',
            '外盘 > 内盘（资金净流入，主动买入 > 主动卖出）',
            'KDJ_K < 80（未超买）',
            'RSI6 < 70（未超买）',
            f'每个板块最多 {config.top_per_sector} 只（0 表示不限制）',
        ],
        start=6,
    ):
        ws.cell(i, 1, f'• {cond}')

    ws['A15'] = '免责声明：本报告仅供参考，不构成投资建议。股市有风险，投资需谨慎。'
    ws['A15'].font = Font(color='808080', size=9)


def _write_picks(wb, stocks):
    ws = wb.create_sheet('精选Top20')
    headers = ['排名', '代码', '名称', '板块', '最新价', '涨幅%', '换手%', '量比', '外盘%', '评分']
    rows = [
        [
            i,
            s['code'],
            s['name'],
            s['sector'],
            round(s['price'], 2),
            round(s['change_pct'], 2),
            round(s['turnover'], 2),
            round(s['volume_ratio'], 2),
            round(s['outer_ratio'], 1),
            s['score'],
        ]
        for i, s in enumerate(stocks[:20], 1)
    ]
    next_row = _write_table(ws, 1, headers, rows, [8, 10, 12, 12, 10, 10, 10, 10, 10, 8])
    ws.cell(next_row + 1, 1, f'共 {len(stocks)} 只（此处展示评分最高的前 20 只）').font = Font(color='808080', size=9)

    if len(stocks) > 20:
        ws2 = wb.create_sheet('全部候选')
        h2 = ['排名', '代码', '名称', '板块', '最新价', '涨幅%', '换手%', '量比', '外盘%', '评分']
        r2 = [
            [
                i,
                s['code'],
                s['name'],
                s['sector'],
                round(s['price'], 2),
                round(s['change_pct'], 2),
                round(s['turnover'], 2),
                round(s['volume_ratio'], 2),
                round(s['outer_ratio'], 1),
                s['score'],
            ]
            for i, s in enumerate(stocks, 1)
        ]
        _write_table(ws2, 1, h2, r2, [8, 10, 12, 12, 10, 10, 10, 10, 10, 8])


def _write_top5(wb, stocks):
    if not stocks:
        return
    ws = wb.create_sheet('Top5分析')
    row = 1
    for i, s in enumerate(stocks[:5], 1):
        ws.cell(row, 1, f'{i}. {s["name"]}（{s["code"]}）— {s["score"]}/80 分').font = TITLE_FONT
        row += 1

        info = [
            ('最新价', f'{s["price"]:.2f} 元'),
            ('涨跌幅', f'{s["change_pct"]:.2f}%'),
            ('换手率', f'{s["turnover"]:.2f}%'),
            ('量比', f'{s["volume_ratio"]:.2f} 倍'),
            ('外盘占比', f'{s["outer_ratio"]:.1f}%（内盘 {s["inner_ratio"]:.1f}%）'),
            ('资金流向', '净流入' if s['outer_ratio'] > 50 else '净流出'),
            ('KDJ_K', f'{s["kdj_k"]:.1f}'),
            ('RSI(6)', f'{s["rsi6"]:.1f}'),
            ('近3日累计涨幅', f'{s["recent_3day_change"]:.2f}%'),
        ]
        for name, val in info:
            ws.cell(row, 1, name).font = SECTION_FONT
            ws.cell(row, 2, val)
            row += 1
        row += 1


def _write_history(wb, stocks):
    if not stocks:
        return
    ws = wb.create_sheet('近5日走势')
    headers = ['名称', '代码', '日期', '收盘价', '涨跌幅%', '换手率%', '成交量(手)']
    rows = []
    for s in stocks[:20]:
        hist = s['hist_data']
        for _, r in hist.iterrows():
            rows.append([
                s['name'],
                s['code'],
                str(r['日期']),
                round(float(r['收盘']), 2),
                round(float(r['涨跌幅']), 2),
                round(float(r['换手率']), 2),
                int(r['成交量']),
            ])
    _write_table(ws, 1, headers, rows, [12, 10, 12, 10, 10, 10, 14])